"""Comprehensive unit and API test suite for Phase 1 (FastAPI backend, extractors, privacy hygiene)."""
import io
import json
import os
import pytest
from fastapi.testclient import TestClient

from aishield.app import app, JOBS, GOV
from aishield.extract import extract_text_from_bytes, DocumentExtractionError, MAX_FILE_SIZE_BYTES, MAX_CHARS

client = TestClient(app)

SAMPLE = (
    "My name is Rindhiya A. My email is rindhiya@example.com, and my register number is 23UAI091. "
    "Reach me on 98765 43210 or 9 8 7 6 5 4 3 2 1 0. Date of birth: 14/03/2004. "
    "Please summarize my project report."
)


# ==================== Extraction Tests ====================

def test_extract_txt_utf8():
    text = "Hello world! Test text file."
    b = text.encode("utf-8")
    assert extract_text_from_bytes(b, "test.txt") == text


def test_extract_txt_latin1():
    text = "Rindhiya résumé café"
    b = text.encode("latin-1")
    assert "café" in extract_text_from_bytes(b, "test.txt") or "cafe" in extract_text_from_bytes(b, "test.txt")


def test_extract_docx():
    import docx
    bio = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("First line of confidential document.")
    doc.add_paragraph("Second line with email test@example.com.")
    doc.save(bio)
    bio.seek(0)
    res = extract_text_from_bytes(bio.read(), "document.docx")
    assert "First line" in res and "test@example.com" in res


def test_extract_pdf():
    import pdfplumber
    # Test valid handling or mock if binary PDF generation requires fpdf
    # We test empty / invalid PDF error handling
    with pytest.raises(DocumentExtractionError):
        extract_text_from_bytes(b"Not a real pdf content", "doc.pdf")


def test_extract_unsupported_extension():
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_text_from_bytes(b"Some text", "image.png")
    assert "Unsupported file type" in str(exc_info.value)


def test_extract_file_too_large():
    huge_bytes = b"x" * (MAX_FILE_SIZE_BYTES + 10)
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_text_from_bytes(huge_bytes, "huge.txt")
    assert exc_info.value.status_code == 413


def test_extract_empty_file():
    with pytest.raises(DocumentExtractionError) as exc_info:
        extract_text_from_bytes(b"   \n  ", "empty.txt")
    assert "empty" in str(exc_info.value).lower()


# ==================== FastAPI Endpoint Tests ====================

def test_api_scan_success():
    resp = client.post("/api/scan", json={"text": SAMPLE, "job_id": "test_job_001"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["job_id"] == "test_job_001"
    assert data["risk"] == "HIGH"
    assert data["entity_count"] >= 4

    # Privacy hygiene: ensure raw sensitive values are NOT returned
    body_str = resp.text
    assert "rindhiya@example.com" not in body_str
    assert "23UAI091" not in body_str
    assert "98765 43210" not in body_str


def test_api_scan_empty_400():
    resp = client.post("/api/scan", json={"text": "   "})
    assert resp.status_code == 400


def test_api_scan_oversized_413():
    resp = client.post("/api/scan", json={"text": "A" * (MAX_CHARS + 1)})
    assert resp.status_code == 413


def test_api_upload_txt():
    txt_content = b"My register number is 23UAI091. Please redact."
    files = {"file": ("student_data.txt", txt_content, "text/plain")}
    resp = client.post("/api/upload", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert any(e["type"] == "STUDENT_ID" for e in data["entities"])
    assert "23UAI091" not in resp.text


def test_api_redact_flow():
    # 1. Scan
    scan_resp = client.post("/api/scan", json={"text": SAMPLE, "job_id": "redact_test_job"})
    assert scan_resp.status_code == 200
    
    # 2. Redact with all approved
    redact_resp = client.post("/api/redact", json={
        "job_id": "redact_test_job",
        "approved_types": ["PERSON", "EMAIL", "STUDENT_ID", "PHONE", "DATE_OF_BIRTH"]
    })
    assert redact_resp.status_code == 200
    r_data = redact_resp.json()
    assert r_data["decision"] == "REDACT"
    assert r_data["residual"] == 0
    assert "[PERSON]" in r_data["sanitized"]
    assert "[EMAIL]" in r_data["sanitized"]
    assert "rindhiya@example.com" not in r_data["sanitized"]


def test_api_redact_request_approval_on_high_risk_kept():
    client.post("/api/scan", json={"text": SAMPLE, "job_id": "approval_test_job"})
    
    # User unselects STUDENT_ID (a high risk type)
    redact_resp = client.post("/api/redact", json={
        "job_id": "approval_test_job",
        "approved_types": ["EMAIL"],
        "override_high": False
    })
    assert redact_resp.status_code == 200
    r_data = redact_resp.json()
    assert r_data["decision"] == "REQUEST_APPROVAL"
    assert r_data["sanitized"] is None


def test_api_redact_override_high():
    client.post("/api/scan", json={"text": SAMPLE, "job_id": "override_test_job"})
    
    # User overrides high risk
    redact_resp = client.post("/api/redact", json={
        "job_id": "override_test_job",
        "approved_types": ["EMAIL"],
        "override_high": True
    })
    assert redact_resp.status_code == 200
    r_data = redact_resp.json()
    assert r_data["decision"] in ("REDACT", "ALLOW", "BLOCK")
    assert r_data["sanitized"] is not None


def test_api_report_endpoint():
    job_id = "report_test_job"
    client.post("/api/scan", json={"text": SAMPLE, "job_id": job_id})
    resp = client.get(f"/api/report/{job_id}")
    assert resp.status_code == 200
    data = resp.json()
    assert data["job_id"] == job_id
    assert data["risk"] == "HIGH"
    assert "EMAIL" in data["entity_counts"]
    assert data["ledger_valid"] is True


def test_api_report_404():
    resp = client.get("/api/report/non_existent_job_xyz")
    assert resp.status_code == 404


def test_api_ledger_endpoint():
    resp = client.get("/api/ledger")
    assert resp.status_code == 200
    data = resp.json()
    assert "is_valid" in data
    assert data["is_valid"] is True
    assert len(data["blocks"]) > 0
    # Ledger blocks must never contain raw sensitive info
    assert "rindhiya@example.com" not in resp.text
    assert "23UAI091" not in resp.text


def test_static_index_page():
    resp = client.get("/")
    assert resp.status_code == 200
    assert "AI Shield" in resp.text
    assert "Rindhiya A" in resp.text
